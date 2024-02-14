import os
import openpyxl
import docx
import PyPDF3

class FileHandler:
    @staticmethod
    def retrieve_and_clean_data(filename):
        # Retrieve and clean the data from files in the local directory
        try:
            file_extension = os.path.splitext(filename)[1].lower()
            if file_extension == '.xlsx':
                return FileHandler.retrieve_excel_data(filename)
            elif file_extension == '.docx':
                return FileHandler.retrieve_docx_data(filename)
            elif file_extension == '.pdf':
                return FileHandler.retrieve_pdf_data(filename)
            elif file_extension == '.py':
                return FileHandler.retrieve_python_data(filename)
            else:
                return FileHandler.retrieve_text_data(filename)
        except Exception as e:
            print(f"An error occurred while opening the file: {e}")
            return None

    @staticmethod
    def retrieve_python_data(filename):
        try:
            with open(filename, 'r', encoding='utf-8', errors='ignore') as file:
                data = file.read()
            return data
        except Exception as e:
            print(f"An error occurred while opening the Python file: {e}")
            return None

    @staticmethod
    def retrieve_excel_data(filename):
        try:
            workbook = openpyxl.load_workbook(filename)
            sheet = workbook.active
            data = []
            for row in sheet.iter_rows(values_only=True):
                data.append(row)
            return data
        except Exception as e:
            print(f"An error occurred while opening the Excel file: {e}")
            return None

    @staticmethod
    def retrieve_docx_data(filename):
        try:
            doc = docx.Document(filename)
            file_contents = ''
            for para in doc.paragraphs:
                file_contents += para.text + '\n'
            return file_contents
        except Exception as e:
            print(f"An error occurred while opening the Word document: {e}")
            return None

    @staticmethod
    def retrieve_pdf_data(filename):
        try:
            with open(filename, 'rb') as file:
                reader = PyPDF3.PdfReader(file)
                file_contents = ''
                for page_num in range(len(reader.pages)):
                    page = reader.pages[page_num]
                    file_contents += page.extract_text()
            return file_contents
        except Exception as e:
            print(f"An error occurred while opening the PDF file: {e}")
            return None

    @staticmethod
    def retrieve_text_data(filename):
        try:
            with open(filename, 'r', encoding='utf-8', errors='ignore') as file:
                data = file.read()
            return data
        except FileNotFoundError:
            return None
        except Exception as e:
            print(f"An error occurred while opening the text file: {e}")
            return None

    @staticmethod
    def get_file_names(directory):
        # Get a list of file names in the specified directory
        file_names = []
        for file_name in os.listdir(directory):
            if os.path.isfile(os.path.join(directory, file_name)):
                file_names.append(file_name)
        return file_names
