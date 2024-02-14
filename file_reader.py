import os
import openai
from llama_index.core import SimpleDirectoryReader, Document
from llama_index.core.readers.base import BaseReader
import pandas as pd
import nest_asyncio
import openpyxl
import textract

openai.api_key = os.environ["OPENAI_API_KEY"]

# Define a custom reader for .txt files
class TxtReader(BaseReader):
    def load_data(self, file, extra_info=None):
        with open(file, "r", encoding='utf-8') as f:
            text = f.read()
        return [Document(text=text, extra_info=extra_info or {})]

# Define a custom reader for .xlsx files using pandas
class XlsxReader(BaseReader):
    def load_data(self, file, extra_info=None):
        dataframe = pd.read_excel(file)
        # Convert the entire dataframe to a single string
        text = dataframe.to_csv(index=False)
        return [Document(text=text, extra_info=extra_info or {})]

# Define a custom reader for .py files (or any plain text files)
class PyReader(BaseReader):
    def load_data(self, file, extra_info=None):
        with open(file, "r", encoding='utf-8') as f:
            text = f.read()
        return [Document(text=text, extra_info=extra_info or {})]

# Instantiate the SimpleDirectoryReader with custom readers for each file extension
reader = SimpleDirectoryReader(
    input_dir="./data",
    file_extractor={
        ".txt": TxtReader(),
        ".xlsx": XlsxReader(),
        ".py": PyReader()
        # Add other file readers as necessary
    }
)

# Load the documents using the reader
documents = reader.load_data()

# Print out the documents
for document in documents:
    print(document.text)

nest_asyncio.apply()
